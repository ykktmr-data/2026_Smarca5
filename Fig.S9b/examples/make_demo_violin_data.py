#!/usr/bin/env python3
"""
make_demo_violin_data.py

Generates a small SYNTHETIC violin_plot_underlying_data.xlsx with the same
layout analyze_signal_overlap.py expects (one sheet per comparison, metadata
rows 1-3, header row 4, signalValue data from row 5 in columns A/B), so the
demo can be run end-to-end without any real ChIP-seq/CUT&Tag data.

This is for demonstrating that the pipeline runs correctly, NOT to reproduce
the numbers in the published figure -- the values here are randomly
generated.

Usage:
    python make_demo_violin_data.py --outdir ./demo_data
"""
import argparse
from pathlib import Path

import numpy as np
import openpyxl

COMPARISONS = [
    ("SMARCA5 vs RAR", "no RAR overlap", "RAR overlap"),
    ("SMARCA5 vs DMRT1", "no DMRT1 overlap", "DMRT1 overlap"),
    ("DMRT1 vs RAR", "no RAR overlap", "RAR overlap"),
    ("DMRT1 vs SMARCA5", "no SMARCA5 overlap", "SMARCA5 overlap"),
]


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--outdir", default="./demo_data")
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--n-per-group", type=int, default=40,
                     help="Number of synthetic signalValue points per group (default: 40)")
    args = ap.parse_args()

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(args.seed)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name, label_a, label_b in COMPARISONS:
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = f"Demo synthetic data for {sheet_name}"
        ws["A2"] = "Randomly generated signalValue distributions (not real data)"
        ws["A3"] = "Columns: A = group without overlap, B = group with overlap"
        ws["A4"] = label_a
        ws["B4"] = label_b

        n_a = args.n_per_group
        n_b = args.n_per_group - rng.integers(0, 5)  # unequal group sizes, as in real data
        group_a = rng.gamma(shape=2.0, scale=5.0, size=n_a)
        group_b = rng.gamma(shape=2.0, scale=9.0, size=n_b)  # shifted up -> a real effect to detect

        for i, v in enumerate(group_a):
            ws.cell(row=5 + i, column=1, value=float(v))
        for i, v in enumerate(group_b):
            ws.cell(row=5 + i, column=2, value=float(v))

    out_path = outdir / "violin_plot_underlying_data.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
